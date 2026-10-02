"""Renders a Result into a native, editable PowerPoint deck (python-pptx).

* real slide titles (title placeholders), native charts and native tables - nothing is a picture,
  so the team can restyle and edit the deck in PowerPoint
* 16:9, Calibri for Latin text and Yu Gothic as the East Asian font for the Japanese category names
* speaker notes on every slide
"""
from __future__ import annotations

import datetime as dt
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .analysis import (DEFEND, DEPRIORITIZE, MONITOR, RECAPTURE, REPLICATE, REVIEW, SCALE, TEST, V_SMALL, Result, _jpy,
                       ORDER, T_DOWN, T_FLAT, T_UP, key_notes)
from .pdf import DIR_BLURB

# palette: deep teal dominant, warm amber as the single sharp accent
DARK, TEAL, ACCENT = "0E3B43", "1C7C7D", "F2A33A"
LIGHT, CARD, TEXT, MUTED = "F3F6F6", "E6EEEE", "1F2933", "5F6C72"
GREEN, RED = "2E7D32", "C62828"
DIR_COLOR = {SCALE: "2E7D32", REPLICATE: "1C7C7D", RECAPTURE: "C77700", TEST: "6A4C93", REVIEW: "B71C1C", DEFEND: "455A64"}
LATIN, EA = "Calibri", "Yu Gothic"
SW, SH = 13.333, 7.5
MX = 0.6  # side margin (inches)


def _rgb(h):
    return RGBColor.from_string(h)


def _style(run, size, color=TEXT, bold=False, italic=False):
    f = run.font
    f.size, f.bold, f.italic = Pt(size), bold, italic
    f.color.rgb = _rgb(color)
    f.name = LATIN
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("a:ea")):
        rPr.remove(old)
    rPr.append(rPr.makeelement(qn("a:ea"), {"typeface": EA}))


def _bullet(paragraph, char="•", indent=0.22):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(bu)


def _text(slide, x, y, w, h, text, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          name=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _style(r, size, color, bold, italic)
    return tb


def _bullets(slide, x, y, w, h, items, size=15, color=TEXT, gap=8, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap, tf.auto_size = True, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        _bullet(p)
        r = p.add_run()
        r.text = it
        _style(r, size, color)
    return tb


def _box(slide, x, y, w, h, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def _title(slide, text, color=DARK):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(MX), Inches(0.38), Inches(SW - 2 * MX), Inches(0.95)
    tf = t.text_frame
    tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    r.text = text
    _style(r, 28, color, bold=True)
    return t


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _table(slide, rows, widths, x, y, row_h=0.42, size=11, right_cols=(), header_fill=DARK, name="Table"):
    n_r, n_c = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * n_r))
    gf.name = name
    tbl = gf.table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i in range(n_r):
        tbl.rows[i].height = Inches(row_h)
        for j in range(n_c):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(header_fill if i == 0 else ("FFFFFF" if i % 2 else LIGHT))
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if j in right_cols else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(rows[i][j])
            _style(r, size, "FFFFFF" if i == 0 else TEXT, bold=(i == 0 or (j == 0 and False)))
    return gf


def _bar_chart(slide, cats, vals, colors, x, y, w, h, name, size=11, label_fmt='#,##0'):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("値", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size, ch.font.name = Pt(size), LATIN
    ch.font.color.rgb = _rgb(TEXT)
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = label_fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.color.rgb = Pt(size), _rgb(MUTED)
    ser = plot.series[0]
    ser.invert_if_negative = False   # otherwise negative bars render white
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(c)
        dpt = ser._element.get_or_add_dPt_for_point(i)   # per-point fill needs its own flag, too
        if dpt.find(qn("c:invertIfNegative")) is None:
            dpt.idx.addnext(dpt.makeelement(qn("c:invertIfNegative"), {"val": "0"}))
    ca, va = ch.category_axis, ch.value_axis
    ca.reverse_order = True
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.fill.background()
    ca.has_major_gridlines = False
    va.visible = False
    va.has_major_gridlines = False
    return gf


def _pct(x, d=1):
    return "－" if x is None or x != x else f"{x * 100:.{d}f}%"


def _spct(x, d=1):
    return "－" if x is None or x != x else f"{x * 100:+.{d}f}%"


def build_pptx(res: Result, source_name: str = "") -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    L_TITLE, L_TONLY = prs.slide_layouts[0], prs.slide_layouts[5]
    prev, cur = res.periods_full
    LAT = res.period_latest                      # the newest period, when it covers less than the others: shown alongside, not compared in yen
    LATH = f"{LAT}期*" if LAT else ""
    LATNOTE = (f"* {LAT}期は他の期より対象期間が短く(売上は{cur}期の{res.kpi['sales_latest'] / res.kpi['sales_cur']:.0%})、"
               "売上は他の期と比べられないため、増減は表示していません。PB比率とPB SKU数は並べて比較できます。") if LAT else ""
    k = res.kpi
    d = res.subs
    prs.core_properties.title = "ペットカテゴリー PB戦略レポート"
    prs.core_properties.author = "PB戦略レポート"

    # ---------------------------------------------------------------- 1. title
    s = prs.slides.add_slide(L_TITLE)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = _rgb(DARK)
    t, sub = s.shapes.title, s.placeholders[1]
    t.left, t.top, t.width, t.height = Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.7)
    t.text_frame.word_wrap = True
    t.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "ペットカテゴリー PB戦略レポート"
    _style(r, 44, "FFFFFF", bold=True)
    sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(3.5), Inches(11.5), Inches(0.9)
    sub.text_frame.word_wrap = True
    sp = sub.text_frame.paragraphs[0]
    sp.alignment = PP_ALIGN.LEFT
    r = sp.add_run()
    extra = f"  |  {LAT}期も並べて表示" if LAT else ""
    r.text = f"{prev}期 vs {cur}期{extra}"
    _style(r, 20, "BFD9DA")
    sp2 = sub.text_frame.add_paragraph()
    sp2.alignment = PP_ALIGN.LEFT
    r = sp2.add_run()
    r.text = f"作成日 {dt.datetime.now():%Y-%m-%d}" + (f"  |  元データ: {source_name}" if source_name else "")
    _style(r, 13, "8FB5B7")
    _notes(s, f"PB戦略資料。比較期間: {prev}期 vs {cur}期。トライアルの方向性の記載は使用していません。")

    # ---------------------------------------------------------------- 2. takeaways
    s = prs.slides.add_slide(L_TONLY)
    _title(s, f"PB比率 {_pct(k['share_prev'])}→{_pct(k['share_cur'])}({(k['share_cur'] - k['share_prev']) * 100:+.2f}pt)、PB SKU数 {_spct(k['sku_growth'], 0)}")
    _bullets(s, MX, 1.7, SW - 2 * MX, 5.2, res.findings, size=20, gap=22, name="Findings")
    _notes(s, "\n".join(res.findings))

    # ---------------------------------------------------------------- 3. PB ratio movers
    s = prs.slides.add_slide(L_TONLY)
    min_s = res.thresholds.min_sales
    mv = d[d.sales >= min_s].sort_values("d_share_pt", ascending=False)
    gains, losses = mv[mv.d_share_pt > 0].head(8), mv[mv.d_share_pt < 0].tail(5).sort_values("d_share_pt")
    sel = list(gains.iterrows()) + list(losses.iterrows())
    _title(s, f"PB比率の増減 {prev}期→{cur}期")
    if sel:
        _bar_chart(s, [r["name"] for _, r in sel], [round(r.d_share_pt, 1) for _, r in sel],
                   [GREEN if r.d_share_pt > 0 else RED for _, r in sel], MX, 1.5, 8.4, 5.4, "PB ratio change chart",
                   label_fmt='+0.0"pt";-0.0"pt"')
    _text(s, MX, 1.28, 8, 0.25, f"PB比率の増減(pt、売上{_jpy(min_s)}以上のサブカテの上位・下位)", 11, MUTED)
    _box(s, 9.3, 1.55, SW - 9.3 - MX, 4.0, name="Mover notes card")
    notes = []
    if len(gains):
        g = gains.iloc[0]
        notes.append(f"最大の上昇: {g['name']}(PB比率 {_pct(g.share_prev)}→{_pct(g.share_cur)}、{g.d_share_pt:+.1f}pt)。")
    if len(losses):
        l = losses.iloc[0]
        notes.append(f"最大の低下: {l['name']}(PB比率 {_pct(l.share_prev)}→{_pct(l.share_cur)}、{l.d_share_pt:+.1f}pt)。")
    notes.append("PB比率は元資料の値です。売上の大きさは考慮せず、比率の変化(pt)だけを並べています。")
    _text(s, 9.55, 1.8, SW - 9.3 - MX - 0.5, 0.4, "グラフの見方", 16, DARK, bold=True)
    _bullets(s, 9.55, 2.35, SW - 9.3 - MX - 0.5, 3.1, notes, size=12, name="Mover notes")
    _notes(s, "棒グラフは2つの通期間のPB比率(元資料の値)の差(pt)です。" + " ".join(notes))

    # ---------------------------------------------------------------- 4. big markets with a low PB ratio
    s = prs.slides.add_slide(L_TONLY)
    ws_ = res.thresholds.white_space_max_share
    pools = d[d.share_cur <= ws_].sort_values("sales", ascending=False).head(10)
    _title(s, "売上の大きい市場にPB比率の低いところがある")
    _text(s, MX, 1.28, 9, 0.25, f"{cur}期の全体売上(百万円、括弧内はPB比率)。PB比率{ws_:.0%}以下のサブカテの上位10", 11, MUTED)
    _bar_chart(s, [f"{r['name']}  (PB {r.share_cur * 100:.1f}%)" for _, r in pools.iterrows()],
               [round(r.sales / 1e6, 1) for _, r in pools.iterrows()], [TEAL] * len(pools),
               MX, 1.5, 8.6, 5.4, "Low PB ratio markets chart")
    _box(s, 9.6, 1.55, SW - 9.6 - MX, 4.0, name="Pool notes card")
    n_low = int((d.share_cur <= ws_).sum())
    pn = [f"PB比率が{ws_:.0%}以下のサブカテは全{len(d)}件中{n_low}件です。売上の大きい順に上位10件を表示しています。",
          "全体売上は元資料の値です。PB売上額は元資料にないため表示していません。",
          "規模だけでは推奨になりません。実現性と粗利はこのデータに含まれていません。"]
    _text(s, 9.85, 1.8, SW - 9.6 - MX - 0.5, 0.4, "示唆", 16, DARK, bold=True)
    _bullets(s, 9.85, 2.35, SW - 9.6 - MX - 0.5, 3.1, pn, size=12, name="Pool notes")
    _notes(s, " ".join(pn))

    # ---------------------------------------------------------------- department comparison
    dp = res.depts
    s = prs.slides.add_slide(L_TONLY)
    _title(s, "部門別の売上とPB比率")
    _text(s, MX, 1.28, 5.2, 0.25, "PB比率の推移(%)", 11, MUTED)
    cd = CategoryChartData()
    cd.categories = [r["dept"] for _, r in dp.iterrows()]
    cd.add_series(f"{prev}期", [round(float(v) * 100, 1) for v in dp.share_prev])
    cd.add_series(f"{cur}期", [round(float(v) * 100, 1) for v in dp.share_cur])
    if LAT:
        cd.add_series(LATH, [round(float(v) * 100, 1) for v in dp.share_latest])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(MX), Inches(1.55), Inches(5.2), Inches(3.9), cd)
    gf.name = "PB ratio by department chart"
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    ch.legend.font.size = Pt(11)
    ch.font.size, ch.font.name = Pt(11), LATIN
    pl = ch.plots[0]
    pl.gap_width, pl.overlap = 50, -5
    pl.has_data_labels = True
    pl.data_labels.number_format, pl.data_labels.number_format_is_linked = '0.0"%"', False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(9)
    for ser, col in zip(pl.series, ("B8C9CA", TEAL, ACCENT)):
        ser.invert_if_negative = False
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb(col)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.reverse_order = True
    ch.category_axis.format.line.fill.background()
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    rows = [["部門", "売上", "市場前年比", "PB比率 前期→当期", "PB SKU数"]]
    for _, r_ in dp.iterrows():
        rows.append([r_["dept"], _jpy(r_.sales), _spct(r_.yoy), f"{_pct(r_.share_prev)}→{_pct(r_.share_cur)}",
                     f"{r_.sku_prev:.0f}→{r_.sku_cur:.0f}"])
    rows.append(["ペット計", _jpy(k["sales_cur"]), _spct(k["sales_yoy"]), f"{_pct(k['share_prev'])}→{_pct(k['share_cur'])}",
                 f"{k['sku_prev']:.0f}→{k['sku_cur']:.0f}"])
    _table(s, rows, [1.75, 1.2, 1.1, 1.5, 1.1], 6.05, 1.55, row_h=0.5, size=11, right_cols=(1, 2, 4), name="Department table")
    _text(s, MX, 5.5, SW - 2 * MX, 0.9, f"ペット全体のPB比率{_pct(k['share_cur'])}は、PB比率の大きく異なる部門の平均です。"
          "部門ごとに状況が違うため、次のスライド以降は部門別に見ていきます。売上とPB比率は元資料の値です。", 13, MUTED)
    if LAT:
        _text(s, MX, 6.45, SW - 2 * MX, 0.6, LATNOTE, 10, MUTED)
    _notes(s, "部門別のPB比率: " + "、".join(f"{r_['dept']} {_pct(r_.share_cur)}" for _, r_ in dp.iterrows()))

    if LAT:
        s = prs.slides.add_slide(L_TONLY)
        _title(s, "3期の推移(部門別)")
        rows = [["部門", f"売上 {prev}", f"売上 {cur}", f"売上 {LATH}", f"PB比率 {prev}", f"PB比率 {cur}", f"PB比率 {LATH}",
                 f"PB SKU数 {prev}", f"PB SKU数 {cur}", f"PB SKU数 {LATH}"]]
        for _, r_ in dp.iterrows():
            rows.append([r_["dept"], _jpy(r_.sales_prev), _jpy(r_.sales), _jpy(r_.sales_latest), _pct(r_.share_prev), _pct(r_.share_cur), _pct(r_.share_latest),
                         f"{r_.sku_prev:.0f}", f"{r_.sku_cur:.0f}", f"{r_.sku_latest:.0f}"])
        rows.append(["ペット計", _jpy(k["sales_prev"]), _jpy(k["sales_cur"]), _jpy(k["sales_latest"]), _pct(k["share_prev"]), _pct(k["share_cur"]), _pct(k["share_latest"]),
                     f"{k['sku_prev']:.0f}", f"{k['sku_cur']:.0f}", f"{k['sku_latest']:.0f}"])
        w_ = [1.9, 1.15, 1.15, 1.25, 1.0, 1.0, 1.1, 1.15, 1.15, 1.25]
        _table(s, rows, w_, MX, 1.55, row_h=0.6, size=10.5, right_cols=tuple(range(1, len(w_))), name="Three-period table")
        _text(s, MX, 1.55 + 0.6 * len(rows) + 0.3, SW - 2 * MX, 0.8, LATNOTE, 11, MUTED)
        _notes(s, LATNOTE + f" {LAT}期: 売上 {_jpy(k['sales_latest'])}、PB比率 {_pct(k['share_latest'])}。")

    # ---------------------------------------------------------------- SKU vs PB ratio
    s = prs.slides.add_slide(L_TONLY)
    t_ = res.thresholds
    cs = res.cats[(res.cats.share_prev >= t_.min_share) & (res.cats.sales_prev >= t_.min_sales) & (res.cats.sku_prev > 0)].sort_values("sku_growth", ascending=False)
    if (k["sku_growth"] or 0) > (k["share_growth"] or 0):
        _title(s, f"PB SKU数{_spct(k['sku_growth'], 0)}に対しPB比率{_pct(k['share_prev'])}→{_pct(k['share_cur'])}: SKU増がPB比率の上昇を上回る")
    else:
        _title(s, f"PB SKU数{_spct(k['sku_growth'], 0)}、PB比率{_pct(k['share_prev'])}→{_pct(k['share_cur'])}: PB比率がSKU増に追随")
    if len(cs):
        cd = CategoryChartData()
        cd.categories = cs["cat"].tolist()
        cd.add_series("PB SKU数の増加率", [round(float(v), 4) for v in cs.sku_growth])
        cd.add_series("PB比率の増加率(相対)", [round(float(v), 4) for v in cs.share_growth])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(MX), Inches(1.75), Inches(7.7), Inches(4.2), cd)
        gf.name = "SKU vs PB ratio growth chart"
        ch = gf.chart
        ch.has_title = False
        ch.has_legend = True
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        ch.legend.font.size = Pt(12)
        ch.font.size, ch.font.name = Pt(12), LATIN
        pl = ch.plots[0]
        pl.gap_width, pl.overlap = 60, -5
        pl.has_data_labels = True
        pl.data_labels.number_format, pl.data_labels.number_format_is_linked = '0%', False
        pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        pl.data_labels.font.size = Pt(11)
        for ser, col in zip(pl.series, (ACCENT, DARK)):
            ser.invert_if_negative = False
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _rgb(col)
        ch.value_axis.visible = False
        ch.value_axis.has_major_gridlines = False
        ch.category_axis.format.line.fill.background()
        ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    _text(s, MX, 1.4, 8, 0.28, f"カテゴリー別の伸び {prev}期→{cur}期(前期PB比率{t_.min_share:.0%}以上、売上{_jpy(t_.min_sales)}以上)", 11, MUTED)
    rows = [["カテゴリー", "PB比率の増減", "判定"]]
    for _, r_ in cs.iterrows():
        rows.append([r_["cat"], f"{r_.d_share_pt:+.1f}pt", r_.verdict])
    _table(s, rows, [1.45, 1.0, 2.05], 8.55, 1.75, row_h=0.38, size=10.5, right_cols=(1,), name="SKU verdict table")
    _notes(s, "カテゴリーごとに2本の棒を比較してください。SKUの棒がPB比率の棒より大きく高い場合、追加したSKUに見合うPB比率の上昇が出ていません。"
              "PB比率の増加率は、前期のPB比率に対する当期のPB比率の変化率(元資料の値から算出)です。")

    # ---------------------------------------------------------------- 5. direction overview
    s = prs.slides.add_slide(L_TONLY)
    _title(s, "サブカテ別の方向性(提案)")
    order = [SCALE, REPLICATE, RECAPTURE, TEST, REVIEW, DEFEND]
    cw, ch_ = (SW - 2 * MX - 0.5) / 3, 2.4
    for i, dn in enumerate(order):
        x, y = MX + (i % 3) * (cw + 0.25), 1.5 + (i // 3) * (ch_ + 0.25)
        x_df = res.directions[dn]
        _box(s, x, y, cw, ch_, name=f"Direction card {i + 1}")
        c = _box(s, x + 0.2, y + 0.2, 0.55, 0.55, fill=DIR_COLOR[dn], shape=MSO_SHAPE.OVAL, name=f"Count {i + 1}")
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(len(x_df))
        _style(r, 16, "FFFFFF", bold=True)
        _text(s, x + 0.95, y + 0.28, cw - 1.15, 0.4, dn, 18, DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _text(s, x + 0.25, y + 0.85, cw - 0.5, 0.85, DIR_BLURB[dn], 10, MUTED)
        tops = "、".join(x_df["name"].head(2).tolist()) if len(x_df) else "現在のルールでは該当なし"
        _text(s, x + 0.25, y + 1.78, cw - 0.5, 0.5, "主な対象: " + tops, 10.5, TEXT, bold=True)
    nd, nm = len(res.directions[DEPRIORITIZE]), len(res.directions[MONITOR])
    _text(s, MX, 6.95, SW - 2 * MX, 0.4, f"このほか、PB未展開で小規模または縮小中の市場が{nd}件({DEPRIORITIZE})、明確なシグナルのないものが{nm}件({MONITOR})。", 11, MUTED)
    _notes(s, f"方向性はPBシェアの水準・推移と、市場の規模・成長に基づいて分類しています。{DEPRIORITIZE}: " +
           "、".join(res.directions[DEPRIORITIZE]["name"].tolist()))

    # ---------------------------------------------------------------- one slide per department
    rank_dir = {key: i for i, key in enumerate(ORDER)}
    for _, dr in dp.iterrows():
        dept = dr["dept"]
        s = prs.slides.add_slide(L_TONLY)
        _title(s, f"{dept}: 売上 {_jpy(dr.sales)}({_spct(dr.yoy, 1)})、PB比率 {_pct(dr.share_cur)}({dr.d_share_pt:+.1f}pt)")
        fl = res.dept_findings.get(dept, [])
        rows = [["カテゴリー", "売上(部門内順位)", "市場成長率", "PB比率(増減pt)", "PB SKU数", "判定"]
                + ([f"{LATH} PB比率/SKU数"] if LAT else [])]
        for _, r_ in res.cats[res.cats["dept"] == dept].sort_values("sales", ascending=False).iterrows():
            rows.append([r_["cat"], f"{_jpy(r_.sales)}({r_.rank_dept}位)", _spct(r_.yoy), f"{_pct(r_.share_cur)}({r_.d_share_pt:+.1f})",
                         f"{r_.sku_prev:.0f}→{r_.sku_cur:.0f}({_spct(r_.sku_growth, 0)})", r_.verdict]
                        + ([f"{_pct(r_.share_latest)}/{r_.sku_latest:.0f}"] if LAT else []))
        cw_ = [1.5, 1.7, 1.1, 1.6, 1.8, 2.9, 1.5] if LAT else [1.7, 2.0, 1.3, 1.9, 2.1, 3.1]
        _table(s, rows, cw_, MX, 1.5, row_h=0.34, size=10, right_cols=(2,), name="Category table")
        y2 = 1.5 + 0.34 * len(rows) + 0.3
        sd = res.subs[res.subs["dept"] == dept].copy()
        sd["_o"] = sd["direction"].map(rank_dir)
        sd = sd.sort_values(["_o", "sales"], ascending=[True, False])
        n_show = max(3, int((SH - 0.55 - y2) / 0.32) - 1)
        show = sd.head(n_show)
        rows = [["サブカテ", "売上", "市場前年比", f"PB比率 {prev}→{cur}" + (f"→{LATH}" if LAT else ""), "PB SKU数"]
                + (["シェア動向"] if LAT else []) + ["方向性"]]
        for _, r_ in show.iterrows():
            sku3 = f"{r_.sku_prev:.0f}→{r_.sku_cur:.0f}" + (f"→{r_.sku_latest:.0f}" if LAT else "")
            lat_ = f"→{r_.share_latest * 100:.1f}%" if LAT and r_.share_latest == r_.share_latest else ""
            trend_ = (f"{ {T_UP: '↑', T_FLAT: '→', T_DOWN: '↓'}[r_.trend] } {r_.mom_pt:+.1f}pt" if r_.trend else "－") if LAT else ""
            rows.append([r_["name"], _jpy(r_.sales), _spct(r_.yoy), f"{r_.share_prev * 100:.1f}%→{r_.share_cur * 100:.1f}%{lat_}", sku3]
                        + ([trend_] if LAT else []) + [r_.direction])
        sw_ = [3.3, 1.2, 1.1, 2.4, 1.2, 1.4, 1.5] if LAT else [3.6, 1.4, 1.3, 2.4, 1.3, 1.7]
        _table(s, rows, sw_, MX, y2, row_h=0.32, size=10, right_cols=(1, 2), name="Sub-category table")
        if len(sd) > len(show):
            _text(s, MX, SH - 0.45, 5, 0.28, f"ほか{len(sd) - len(show)}件(PDFレポート参照)", 11, MUTED)
        if LAT:
            _text(s, MX + (3.2 if len(sd) > len(show) else 0), SH - 0.5, 9.0, 0.4,
                  LATNOTE + f" シェア動向は{LAT}期と{cur}期のPB比率の差(±{res.thresholds.momentum_pt:g}pt以上で上昇・低下)。", 8.5, MUTED)
        _notes(s, "\n".join(fl))

    # ---------------------------------------------------------------- basis
    s = prs.slides.add_slide(L_TONLY)
    _title(s, "前提・注意事項")
    notes_ = key_notes(res)
    _bullets(s, MX, 1.6, SW - 2 * MX, 5.4, notes_, size=18, gap=16, name="Notes")
    _notes(s, "\n".join(notes_))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
